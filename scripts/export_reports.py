import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import reportlab
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def parse_markdown(md_text):
    lines = md_text.splitlines()
    blocks = []
    i = 0
    n = len(lines)
    
    while i < n:
        line = lines[i]
        
        # Heading
        m_head = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m_head:
            level = len(m_head.group(1))
            text = m_head.group(2).strip()
            blocks.append({'type': 'heading', 'level': level, 'text': text})
            i += 1
            continue
            
        # Horizontal rule
        if re.match(r'^(\*{3,}|-{3,}|_{3,})$', line.strip()):
            blocks.append({'type': 'hr'})
            i += 1
            continue
            
        # Code block
        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            if i < n:
                i += 1
            blocks.append({'type': 'code', 'text': '\n'.join(code_lines)})
            continue
            
        # Table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # parse table rows
            rows = []
            for tline in table_lines:
                # check if separator
                if re.match(r'^\|(?:\s*[:-]+[-| :]*)\|?$', tline):
                    continue
                parts = [c.strip() for c in tline.split('|')[1:-1]]
                rows.append(parts)
            blocks.append({'type': 'table', 'rows': rows})
            continue
            
        # Bullet list
        m_bullet = re.match(r'^([ \t]*)[-*+]\s+(.*)$', line)
        if m_bullet:
            indent = len(m_bullet.group(1))
            text = m_bullet.group(2).strip()
            blocks.append({'type': 'bullet', 'text': text, 'indent': indent})
            i += 1
            continue
            
        # Numbered list
        m_num = re.match(r'^([ \t]*)\d+\.\s+(.*)$', line)
        if m_num:
            indent = len(m_num.group(1))
            text = m_num.group(2).strip()
            blocks.append({'type': 'numbered', 'text': text, 'indent': indent})
            i += 1
            continue
            
        # Blockquote
        if line.startswith('>'):
            quote_lines = []
            while i < n and (lines[i].startswith('>') or (lines[i].strip() and not lines[i].startswith('#'))):
                clean_q = re.sub(r'^>\s?', '', lines[i])
                quote_lines.append(clean_q)
                i += 1
            blocks.append({'type': 'blockquote', 'text': '\n'.join(quote_lines)})
            continue
            
        # Paragraph
        if line.strip():
            p_lines = [line.strip()]
            i += 1
            while i < n and lines[i].strip() and not lines[i].startswith('#') and not lines[i].startswith('|') and not lines[i].startswith('```') and not lines[i].startswith('---') and not re.match(r'^[-*+]\s+', lines[i].strip()) and not re.match(r'^\d+\.\s+', lines[i].strip()) and not lines[i].startswith('>'):
                p_lines.append(lines[i].strip())
                i += 1
            blocks.append({'type': 'paragraph', 'text': ' '.join(p_lines)})
            continue
            
        i += 1
        
    return blocks


def clean_markdown_inline(text):
    # Remove images
    text = re.sub(r'!\[.*?\]\(.*?\)', '', text)
    # Convert markdown links [text](url) to text
    text = re.sub(r'\[(.*?)\]\((.*?)\)', r'\1', text)
    # Clean LaTeX math
    text = text.replace(r'$\ge 18$', '>= 18')
    text = text.replace(r'($\ge 18$)', '(>= 18)')
    text = text.replace(r'$\ge 12$', '>= 12')
    text = text.replace(r'$\ge 8$', '>= 8')
    text = text.replace(r'$\le 200\text{ms}$', '<= 200ms')
    text = text.replace(r'($\le 200\text{ms}$)', '(<= 200ms)')
    text = text.replace(r'($T_{\text{select}}$)', '(T_select)')
    text = text.replace(r'$T_{\text{select}}$', 'T_select')
    text = text.replace(r'$V_{\text{select}}$', 'V_select')
    text = text.replace(r'$C_{\text{score}}$', 'C_score')
    text = text.replace(r'$C_{\text{inspect}}$', 'C_inspect')
    text = text.replace(r'$R_{\text{complete}}$', 'R_complete')
    text = text.replace(r'$H_1$', 'H1')
    text = text.replace(r'$H_{1A}$', 'H1A')
    text = text.replace(r'$H_{1B}$', 'H1B')
    text = text.replace(r'$H_2$', 'H2')
    text = text.replace(r'$H_3$', 'H3')
    text = text.replace(r'$p < 0.05$', 'p < 0.05')
    text = text.replace(r'$N = 10\text{--}20$', 'N = 10-20')
    text = text.replace(r'$S = 0.40T + 0.35G + 0.25L$', 'S = 0.40*T + 0.35*G + 0.25*L')
    text = text.replace(r'($S = 0.40T + 0.35G + 0.25L$)', '(S = 0.40*T + 0.35*G + 0.25*L)')
    text = text.replace(r'($0.40T + 0.35G + 0.25L$)', '(0.40*T + 0.35*G + 0.25*L)')
    text = text.replace(r'$0.40T + 0.35G + 0.25L$', '0.40*T + 0.35*G + 0.25*L')
    text = text.replace(r'($0.40T+0.35G+0.25L$)', '(0.40*T + 0.35*G + 0.25*L)')
    text = text.replace(r'$0.40T+0.35G+0.25L$', '0.40*T + 0.35*G + 0.25*L')
    text = text.replace(r'($H_2 \ge 80\%$)', '(H2 >= 80%)')
    text = text.replace(r'($H_3 > 68.0$)', '(H3 > 68.0)')
    text = text.replace(r'$$\text{users (UUID ASC)} \to \text{study\_listings (ID ASC)} \to \text{listing\_responses (ID ASC)} \to \text{study\_sessions (ID ASC)} \to \text{session\_outcomes (ID ASC)}$$', 'users (UUID ASC) -> study_listings (ID ASC) -> listing_responses (ID ASC) -> study_sessions (ID ASC) -> session_outcomes (ID ASC)')
    text = text.replace(r'$$\text{Score} = 100 \times \left(0.40 \cdot T(\text{tags}) + 0.35 \cdot G(\text{goal}) + 0.25 \cdot L(\text{level})\right)$$', 'Score = 100 * (0.40*T(tags) + 0.35*G(goal) + 0.25*L(level))')
    return text


def build_docx(blocks, output_path):
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    for b in blocks:
        b_type = b['type']
        
        if b_type == 'heading':
            level = b['level']
            text = clean_markdown_inline(b['text'])
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            
            if level == 1:
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(text)
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x11, 0x33, 0x66)
            elif level == 2:
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(text)
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            elif level == 3:
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(text)
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            else:
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(text)
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                
        elif b_type == 'paragraph':
            text = clean_markdown_inline(b['text'])
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            # format bold/italics
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    r = p.add_run(part[1:-1])
                    r.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    r = p.add_run(part[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0x80, 0x00, 0x20)
                else:
                    p.add_run(part)
                    
        elif b_type == 'bullet':
            text = clean_markdown_inline(b['text'])
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    r = p.add_run(part[1:-1])
                    r.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    r = p.add_run(part[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                else:
                    p.add_run(part)
                    
        elif b_type == 'numbered':
            text = clean_markdown_inline(b['text'])
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    r = p.add_run(part[1:-1])
                    r.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    r = p.add_run(part[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                else:
                    p.add_run(part)
                    
        elif b_type == 'blockquote':
            text = clean_markdown_inline(b['text'])
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            
        elif b_type == 'code':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(b['text'])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
            
        elif b_type == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # bottom border
            pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                             r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
                             r'</w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            
        elif b_type == 'table':
            rows = b['rows']
            if not rows:
                continue
            num_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            
            for r_idx, row_data in enumerate(rows):
                for c_idx in range(num_cols):
                    cell = table.cell(r_idx, c_idx)
                    val = clean_markdown_inline(row_data[c_idx]) if c_idx < len(row_data) else ''
                    cell.text = val
                    
                    # Style cell text
                    cell_p = cell.paragraphs[0]
                    cell_p.paragraph_format.space_before = Pt(2)
                    cell_p.paragraph_format.space_after = Pt(2)
                    for r in cell_p.runs:
                        r.font.size = Pt(8.5)
                        if r_idx == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            
                    # Cell shading for header
                    if r_idx == 0:
                        shading = parse_xml(r'<w:shd {} w:fill="1F4E79"/>'.format(nsdecls('w')))
                        cell._tc.get_or_add_tcPr().append(shading)
                    elif r_idx % 2 == 1:
                        shading = parse_xml(r'<w:shd {} w:fill="F2F5F8"/>'.format(nsdecls('w')))
                        cell._tc.get_or_add_tcPr().append(shading)
                        
            doc.add_paragraph() # space after table

    doc.save(output_path)
    print(f'DOCX saved to {output_path}')


def build_pdf(blocks, output_path):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    # Register Arial for full Unicode & Cyrillic glyph coverage
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.pdfmetrics import registerFontFamily

    has_arial = False
    if os.path.exists('C:/Windows/Fonts/arial.ttf'):
        pdfmetrics.registerFont(TTFont('Arial', 'C:/Windows/Fonts/arial.ttf'))
        pdfmetrics.registerFont(TTFont('Arial-Bold', 'C:/Windows/Fonts/arialbd.ttf'))
        pdfmetrics.registerFont(TTFont('Arial-Italic', 'C:/Windows/Fonts/ariali.ttf'))
        pdfmetrics.registerFont(TTFont('Arial-BoldItalic', 'C:/Windows/Fonts/arialbi.ttf'))
        registerFontFamily('Arial', normal='Arial', bold='Arial-Bold', italic='Arial-Italic', boldItalic='Arial-BoldItalic')
        has_arial = True

    font_normal = 'Arial' if has_arial else 'Helvetica'
    font_bold = 'Arial-Bold' if has_arial else 'Helvetica-Bold'
    font_italic = 'Arial-Italic' if has_arial else 'Helvetica-Oblique'

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName=font_bold,
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#113366'),
        spaceAfter=12,
        keepWithNext=True
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName=font_bold,
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1F4E79'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )
    h3_style = ParagraphStyle(
        'DocH3',
        parent=styles['Heading3'],
        fontName=font_bold,
        fontSize=10.5,
        leading=13,
        textColor=colors.HexColor('#2E75B6'),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName=font_normal,
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor('#222222'),
        spaceAfter=5
    )
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=body_style,
        fontName=font_normal,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=3
    )
    code_style = ParagraphStyle(
        'DocCode',
        parent=styles['Code'],
        fontSize=7,
        leading=8.5,
        fontName='Courier',
        textColor=colors.HexColor('#111111'),
        backColor=colors.HexColor('#F4F6F9'),
        borderPadding=5,
        spaceBefore=4,
        spaceAfter=6
    )
    quote_style = ParagraphStyle(
        'DocQuote',
        parent=body_style,
        fontName=font_italic,
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#555555'),
        leftIndent=20,
        rightIndent=20,
        spaceBefore=4,
        spaceAfter=6
    )
    th_style = ParagraphStyle(
        'TH',
        fontName=font_bold,
        fontSize=7.5,
        leading=9,
        textColor=colors.white,
        alignment=0
    )
    td_style = ParagraphStyle(
        'TD',
        fontName=font_normal,
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor('#222222'),
        alignment=0
    )
    
    story = []
    
    def format_html_text(text):
        text = clean_markdown_inline(text)
        # escape & < >
        text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        # bold
        text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
        # italics
        text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', text)
        # code
        text = re.sub(r'`(.*?)`', r'<font name="Courier" color="#800020">\1</font>', text)
        return text

    for b in blocks:
        b_type = b['type']
        
        if b_type == 'heading':
            level = b['level']
            html_text = format_html_text(b['text'])
            if level == 1:
                story.append(Paragraph(html_text, title_style))
            elif level == 2:
                story.append(Paragraph(html_text, h2_style))
            elif level == 3:
                story.append(Paragraph(html_text, h3_style))
            else:
                story.append(Paragraph(f'<b>{html_text}</b>', body_style))
                
        elif b_type == 'paragraph':
            html_text = format_html_text(b['text'])
            story.append(Paragraph(html_text, body_style))
            
        elif b_type == 'bullet':
            html_text = '&bull; ' + format_html_text(b['text'])
            story.append(Paragraph(html_text, bullet_style))
            
        elif b_type == 'numbered':
            html_text = '&diams; ' + format_html_text(b['text'])
            story.append(Paragraph(html_text, bullet_style))
            
        elif b_type == 'blockquote':
            html_text = format_html_text(b['text']).replace('\n', '<br/>')
            story.append(Paragraph(f'<i>{html_text}</i>', quote_style))
            
        elif b_type == 'code':
            raw = b['text'].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            raw = raw.replace('\n', '<br/>')
            story.append(Paragraph(raw, code_style))
            
        elif b_type == 'hr':
            story.append(HRFlowable(width='100%', thickness=0.8, color=colors.HexColor('#CCCCCC'), spaceBefore=6, spaceAfter=8))
            
        elif b_type == 'table':
            rows = b['rows']
            if not rows:
                continue
            num_cols = max(len(r) for r in rows)
            
            # Estimate column widths
            total_width = 532 # 612 letter - 80 margins
            col_width = total_width / num_cols
            
            table_data = []
            for r_idx, row in enumerate(rows):
                row_data = []
                for c_idx in range(num_cols):
                    raw_val = row[c_idx] if c_idx < len(row) else ''
                    html_val = format_html_text(raw_val)
                    if r_idx == 0:
                        p = Paragraph(html_val, th_style)
                    else:
                        p = Paragraph(html_val, td_style)
                    row_data.append(p)
                table_data.append(row_data)
                
            t = Table(table_data, colWidths=[col_width]*num_cols, repeatRows=1)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E79')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('LEFTPADDING', (0, 0), (-1, -1), 3),
                ('RIGHTPADDING', (0, 0), (-1, -1), 3),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#DDDDDD')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
            ]))
            story.append(Spacer(1, 4))
            story.append(t)
            story.append(Spacer(1, 6))

    doc.build(story)
    print(f'PDF saved to {output_path}')


if __name__ == '__main__':
    md_path = os.path.abspath('docs/supervisor-review-report.md')
    docx_path = os.path.abspath('docs/supervisor-review-report.docx')
    pdf_path = os.path.abspath('docs/supervisor-review-report.pdf')
    
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
        
    blocks = parse_markdown(md_text)
    print(f'Parsed {len(blocks)} blocks from {md_path}')
    
    build_docx(blocks, docx_path)
    build_pdf(blocks, pdf_path)
